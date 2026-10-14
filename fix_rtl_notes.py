import re, sys
from pptx import Presentation

# pptxgenjs writes each note as one paragraph with raw newlines. Split it into real
# paragraphs, then mark the Arabic ones right-to-left so only they are right-aligned.
path = sys.argv[1]
prs = Presentation(path)
arabic = re.compile(r"[؀-ۿ]")
rtl = 0
for slide in prs.slides:
    if not slide.has_notes_slide:
        continue
    tf = slide.notes_slide.notes_text_frame
    full = tf.text
    tf.text = full
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs)
        if arabic.search(text):
            pPr = para._p.get_or_add_pPr()
            pPr.set("rtl", "1")
            pPr.set("algn", "r")
            rtl += 1
prs.save(path)
print("rtl paragraphs:", rtl)
